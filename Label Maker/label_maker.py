"""Zach Store Label Maker — core.

Reads the inventory workbook, looks up specs and a photo for each item, and writes a
print-ready PDF: an item list followed by cut-out shelf labels.
"""
from __future__ import annotations

import hashlib
import io
import json
import os
import re
import time
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

import openpyxl
import requests
from PIL import Image as PILImage
from reportlab.lib import colors
from reportlab.lib.enums import TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (Image, KeepInFrame, PageBreak, Paragraph, SimpleDocTemplate,
                                Spacer, Table, TableStyle)

TOOL_DIR = Path(__file__).resolve().parent
STORE_DIR = TOOL_DIR.parent
IMAGES_DIR = TOOL_DIR / "images"            # drop your own photos here, named after the item
CACHE_DIR = IMAGES_DIR / "web-cache"        # downloaded web photos
CACHE_INDEX = CACHE_DIR / "index.json"
LOGO = STORE_DIR / "assets" / "zach-logo-mark.png"

DEFAULT_INVENTORY = STORE_DIR / "Zach_Store_Inventory.xlsx"
SAMPLE_INVENTORY = STORE_DIR / "sample-data" / "Sample_Inventory.xlsx"   # demo data shipped with the repo


def default_inventory() -> Path:
    """The real inventory if it's here, otherwise the sample (e.g. in a fresh clone)."""
    return DEFAULT_INVENTORY if DEFAULT_INVENTORY.exists() else SAMPLE_INVENTORY
SPECS_NAME = "Item Specs.xlsx"
PDF_NAME = "Zach Store Item Labels.pdf"

USER_AGENT = "ZachStoreLabelMaker/1.0 (local inventory label tool)"
MISS_RETRY_SECONDS = 24 * 3600              # look again for a web photo once a day
LOW_STOCK = 5

INK = colors.HexColor("#201e1d")
ACCENT = colors.HexColor("#ec3013")
ACCENT_DARK = colors.HexColor("#ae1800")
MUTED = colors.HexColor("#605d5d")
RULE = colors.HexColor("#bdbaba")
PAPER = colors.HexColor("#f3f2f2")

# Photos already in the store website's assets folder.
LOCAL_PHOTOS = {
    "musony w1": "musony-w1.jpg",
    "redmi note 13": "redmi-note-13.jpg",
    "hp 15s": "hp-15s.jpg",
}

# Starter specs for the items in the inventory, written into Item Specs.xlsx the first time.
STARTER_SPECS = {
    "Redmi Note 13 5G": ["6.67\" AMOLED, 120 Hz", "MediaTek Dimensity 6080", "6 GB RAM, 128 GB storage", "108 MP main camera", "5,000 mAh battery, 33 W charging"],
    "Samsung Galaxy M35": ["6.6\" Super AMOLED, 120 Hz", "Exynos 1380", "6 GB RAM, 128 GB storage", "50 MP main camera", "6,000 mAh battery"],
    "iPhone 15 (128GB)": ["6.1\" Super Retina XDR OLED", "A16 Bionic chip", "128 GB storage", "48 MP main camera", "USB-C, Dynamic Island"],
    "MUSONY W1": ["16\" 2.5K IPS, 120 Hz", "Intel Core Ultra 7 + RTX 4090", "120 GB RAM", "99 Wh battery, up to 15 hr", "Storage: TBC"],
    "HP 15s Ryzen 5": ["15.6\" Full HD display", "AMD Ryzen 5 5500U, 6-core", "16 GB RAM", "512 GB SSD", "Up to 9 hr battery"],
    "Lenovo IdeaPad Slim 3": ["15.6\" Full HD IPS display", "AMD Ryzen 5 7520U", "8 GB RAM", "512 GB SSD", "Up to 10 hr battery"],
    "boAt Airdopes 141": ["True wireless earbuds", "Up to 42 hr total playback", "ENx noise cancellation for calls", "BEAST mode low-latency gaming", "IPX4 water resistance"],
    "Anker 20W Fast Charger": ["20 W USB-C Power Delivery", "Fast-charges iPhone and Android", "Compact wall plug", "Built-in safety protection"],
    "Logitech M331 Wireless Mouse": ["SilentTouch quiet clicks", "2.4 GHz USB receiver", "Up to 24-month battery (1 x AA)", "Right-handed contoured shape"],
    "Tempered Glass Screen Guard": ["9H hardness tempered glass", "Oleophobic anti-fingerprint coating", "Bubble-free installation", "Case-friendly fit"],
}

# Generic photo used when no exact product photo is found, chosen by keywords in the name or type.
# Each value is a Wikipedia article whose lead photo suits that kind of product.
GENERIC_TOPICS = [
    (("airdopes", "airpods", "earbud", "buds", "tws", "earphone"), "AirPods"),
    (("headphone", "headset"), "Headphones"),
    (("charger", "adapter", "adaptor"), "AC adapter"),
    (("mouse",), "Computer mouse"),
    (("keyboard",), "Computer keyboard"),
    (("screen guard", "screen protector", "tempered glass"), "Screen protector"),
    (("watch",), "Smartwatch"),
    (("laptop", "notebook", "ideapad", "thinkpad", "vivobook", "macbook"), "Laptop"),
    (("phone", "iphone", "galaxy", "redmi", "pixel", "oneplus", "vivo", "oppo", "realme"), "Smartphone"),
]
NOISE_TOKENS = {"5g", "4g", "lte", "wireless", "fast", "the", "with", "and", "new"}
VARIANT_WORDS = {"pro", "max", "plus", "ultra", "mini", "lite", "fe", "neo", "air", "prime", "edge", "fold", "flip", "se", "slim", "gaming"}


# ── Fonts ───────────────────────────────────────────────────────────────────
def _register_fonts() -> tuple[str, str]:
    fonts = Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts"
    for regular, bold in (("segoeui.ttf", "segoeuib.ttf"), ("arial.ttf", "arialbd.ttf")):
        if (fonts / regular).exists() and (fonts / bold).exists():
            pdfmetrics.registerFont(TTFont("Body", str(fonts / regular)))
            pdfmetrics.registerFont(TTFont("Body-Bold", str(fonts / bold)))
            return "Body", "Body-Bold"
    return "Helvetica", "Helvetica-Bold"


FONT, FONT_BOLD = _register_fonts()
RUPEE = "\u20b9" if FONT == "Body" else "Rs. "


# ── Data ────────────────────────────────────────────────────────────────────
@dataclass
class Item:
    name: str
    type: str = ""
    price: float | None = None
    stock: float | None = None
    specs: list[str] = field(default_factory=list)
    image_hint: str = ""                 # path or URL from the sheets, if given
    photo: Path | None = None
    photo_note: str = ""                 # e.g. "Representative photo · Wikimedia Commons"


def _norm(text) -> str:
    return re.sub(r"\s+", " ", str(text or "")).strip()


def _key(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", name.lower()).strip()


def _num(value):
    if isinstance(value, (int, float)):
        return value
    try:
        return float(str(value).replace(",", "").replace(RUPEE, "").strip())
    except ValueError:
        return None


def _split_specs(value) -> list[str]:
    parts = re.split(r"[\n;•]+", str(value or ""))
    return [p.strip(" -\t") for p in parts if p.strip(" -\t")]


def _find_header(ws, must: str = "model"):
    """Return (row index, {header lower: column index}) for the first row containing `must`."""
    for r, row in enumerate(ws.iter_rows(min_row=1, max_row=30, values_only=True), start=1):
        heads = {_norm(v).lower(): c for c, v in enumerate(row, start=1) if v is not None}
        if any(must in h for h in heads):
            return r, heads
    return None, {}


def _col(heads: dict, *names):
    for want in names:
        for head, col in heads.items():
            if head.startswith(want):
                return col
    return None


def read_inventory(path: Path) -> list[Item]:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb.worksheets[0]
        header_row, heads = _find_header(ws)
        if not header_row:
            raise ValueError("Couldn't find a 'Model Name' header row in the inventory.")
        c_name = _col(heads, "model", "name", "item")
        c_type = _col(heads, "type", "category")
        c_price = _col(heads, "sell", "price", "mrp")
        c_stock = _col(heads, "stock", "qty", "quantity")
        c_specs = _col(heads, "spec")
        c_image = _col(heads, "image", "photo", "picture")
        items, seen = [], set()
        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            get = lambda c: row[c - 1] if c and c - 1 < len(row) else None
            name = _norm(get(c_name))
            if not name:
                continue
            low = name.lower()
            if low in ("total", "how to use") or name.startswith("•"):
                break
            if _key(name) in seen:
                continue
            seen.add(_key(name))
            items.append(Item(name=name, type=_norm(get(c_type)), price=_num(get(c_price)),
                              stock=_num(get(c_stock)), specs=_split_specs(get(c_specs)),
                              image_hint=_norm(get(c_image))))
        return items
    finally:
        wb.close()


# ── Item Specs.xlsx (companion sheet the user edits) ────────────────────────
def specs_path_for(inventory: Path) -> Path:
    return inventory.with_name(SPECS_NAME)


def read_specs(path: Path) -> dict[str, tuple[list[str], str]]:
    if not path.exists():
        return {}
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb.worksheets[0]
        header_row, heads = _find_header(ws)
        c_name, c_specs, c_image = _col(heads, "model", "name"), _col(heads, "spec"), _col(heads, "image", "photo")
        out = {}
        for row in ws.iter_rows(min_row=(header_row or 1) + 1, values_only=True):
            get = lambda c: row[c - 1] if c and c - 1 < len(row) else None
            name = _norm(get(c_name))
            if name:
                out[_key(name)] = (_split_specs(get(c_specs)), _norm(get(c_image)))
        return out
    finally:
        wb.close()


def sync_specs_file(path: Path, items: list[Item], log) -> bool:
    """Create the specs sheet or append rows for new items. Returns True if the file was written."""
    from openpyxl.styles import Alignment, Font, PatternFill

    existing = read_specs(path)
    missing = [it for it in items if _key(it.name) not in existing]
    if path.exists() and not missing:
        return False
    if path.exists():
        wb = openpyxl.load_workbook(path)
        ws = wb.worksheets[0]
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Specs"
        ws.append(["Model Name", "Specs (one per line)", "Image (optional: file path or web link)"])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="201E1D")
        ws.column_dimensions["A"].width = 32
        ws.column_dimensions["B"].width = 60
        ws.column_dimensions["C"].width = 48
        ws.freeze_panes = "A2"
    starter = {_key(k): v for k, v in STARTER_SPECS.items()}
    for it in missing:
        specs = starter.get(_key(it.name), [])
        ws.append([it.name, "\n".join(specs), ""])
        ws.cell(ws.max_row, 2).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(ws.max_row, 1).alignment = Alignment(vertical="top")
    try:
        wb.save(path)
    except PermissionError:
        log(f"Couldn't add new items to {path.name} — it's open in Excel. Will try again later.")
        return False
    names = ", ".join(it.name for it in missing)
    log(f"Added to {path.name}: {names}" if existing else f"Created {path.name} with {len(missing)} items.")
    return True


# ── Photos ──────────────────────────────────────────────────────────────────
def _safe_filename(name: str) -> str:
    return re.sub(r'[<>:"/\\|?*]+', "_", name).strip()


def _load_cache() -> dict:
    try:
        return json.loads(CACHE_INDEX.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}


def _save_cache(cache: dict):
    CACHE_DIR.mkdir(parents=True, exist_ok=True)
    CACHE_INDEX.write_text(json.dumps(cache, indent=2), encoding="utf-8")


def _session() -> requests.Session:
    s = requests.Session()
    s.headers["User-Agent"] = USER_AGENT
    return s


def _polite_get(session, url, tries=4, **kw):
    """GET with a short pause between calls and back-off when Wikimedia says "slow down" (HTTP 429)."""
    for attempt in range(tries):
        time.sleep(0.6)
        r = session.get(url, timeout=20, **kw)
        if r.status_code == 429 and attempt < tries - 1:
            retry_after = r.headers.get("Retry-After", "")
            time.sleep(min(int(retry_after), 20) if retry_after.isdigit() else 3 * (attempt + 1))
            continue
        r.raise_for_status()
        return r
    return r


def _get_json(session, url, params):
    return _polite_get(session, url, params=params).json()


def _download_image(session, url: str, dest_stem: Path) -> Path | None:
    """Download, verify it's a real image, and save as a JPEG (max 900 px)."""
    r = _polite_get(session, url)
    img = PILImage.open(io.BytesIO(r.content))
    img.load()
    if img.mode in ("RGBA", "LA", "P"):
        img = img.convert("RGBA")
        bg = PILImage.new("RGB", img.size, (255, 255, 255))
        bg.paste(img, mask=img.split()[-1])
        img = bg
    img = img.convert("RGB")
    img.thumbnail((900, 900))
    dest = dest_stem.with_suffix(".jpg")
    dest.parent.mkdir(parents=True, exist_ok=True)
    img.save(dest, "JPEG", quality=85)
    return dest


def _tokens(text: str) -> list[str]:
    return [t for t in _key(re.sub(r"\(.*?\)", " ", text)).split() if t not in NOISE_TOKENS]


def _title_matches(name: str, title: str) -> bool:
    want, have = _tokens(name), set(_tokens(title))
    if not want:
        return False
    # "iPhone 15 Pro" or "AirPods Max" is a different product from "iPhone 15" / "AirPods 4".
    if (have & VARIANT_WORDS) - set(want):
        return False
    numeric = [t for t in want if any(ch.isdigit() for ch in t)]
    if any(t not in have for t in numeric):
        return False
    return sum(t in have for t in want) / len(want) >= 0.75


def _wikipedia_photo(session, name: str) -> tuple[str, str] | None:
    """Lead photo of a Wikipedia article whose title matches the product name."""
    data = _get_json(session, "https://en.wikipedia.org/w/api.php", {
        "action": "query", "format": "json", "generator": "search", "gsrsearch": re.sub(r"\(.*?\)", " ", name),
        "gsrlimit": 5, "prop": "pageimages", "piprop": "thumbnail", "pithumbsize": 800})
    pages = sorted(data.get("query", {}).get("pages", {}).values(), key=lambda p: p.get("index", 99))
    for p in pages:
        src = p.get("thumbnail", {}).get("source")
        if src and _title_matches(name, p["title"]):
            return src, "Wikipedia"
    return None


def _topic_photo(session, topic: str) -> tuple[str, str] | None:
    """Lead photo of a general Wikipedia article (e.g. "Computer mouse")."""
    data = _get_json(session, "https://en.wikipedia.org/w/api.php", {
        "action": "query", "format": "json", "titles": topic, "redirects": 1,
        "prop": "pageimages", "piprop": "thumbnail", "pithumbsize": 800})
    for p in data.get("query", {}).get("pages", {}).values():
        src = p.get("thumbnail", {}).get("source")
        if src:
            return src, "Wikipedia"
    return None


def _commons_photo(session, name: str) -> tuple[str, str] | None:
    query = re.sub(r"\(.*?\)", " ", name)
    data = _get_json(session, "https://commons.wikimedia.org/w/api.php", {
        "action": "query", "format": "json", "generator": "search", "gsrsearch": query + " filetype:bitmap",
        "gsrnamespace": 6, "gsrlimit": 5, "prop": "imageinfo", "iiprop": "url", "iiurlwidth": 800})
    pages = sorted(data.get("query", {}).get("pages", {}).values(), key=lambda p: p.get("index", 99))
    for p in pages:
        title = re.sub(r"^File:|\.\w+$", "", p["title"])
        info = (p.get("imageinfo") or [{}])[0]
        if _title_matches(name, title) and info.get("thumburl"):
            return info["thumburl"], "Wikimedia Commons"
    return None


def _generic_topic(item: Item) -> str | None:
    hay = f"{item.name} {item.type}".lower()
    for words, topic in GENERIC_TOPICS:
        if any(w in hay for w in words):
            return topic
    return None


def find_photo(item: Item, session, cache: dict, web: dict, log) -> tuple[Path | None, str]:
    stem = _safe_filename(item.name)

    # 1. Your own photo in Label Maker\images, named like the item.
    for ext in (".jpg", ".jpeg", ".png", ".webp"):
        p = IMAGES_DIR / f"{stem}{ext}"
        if p.exists():
            return p, ""
    lower = {f.stem.lower(): f for f in IMAGES_DIR.glob("*") if f.is_file()}
    if stem.lower() in lower:
        return lower[stem.lower()], ""

    # 2. Image column: a local path or a web link.
    hint = item.image_hint
    if hint:
        local = Path(hint).expanduser()
        if not local.is_absolute():
            local = STORE_DIR / local
        if local.exists():
            return local, ""
        if hint.lower().startswith(("http://", "https://")):
            key = "url:" + hint
            dest = CACHE_DIR / ("link-" + hashlib.sha1(hint.encode()).hexdigest()[:12])
            hit = cache.get(key)
            if hit and Path(hit["file"]).exists():
                return Path(hit["file"]), ""
            try:
                path = _download_image(session, hint, dest)
                cache[key] = {"file": str(path), "note": ""}
                return path, ""
            except Exception as e:  # noqa: BLE001 — any bad link just falls through
                log(f"Couldn't use the image link for {item.name}: {e}")

    # 3. Photos from the store website.
    k = _key(item.name)
    for prefix, filename in LOCAL_PHOTOS.items():
        if k.startswith(prefix) and (STORE_DIR / "assets" / filename).exists():
            return STORE_DIR / "assets" / filename, ""

    # 4. Exact product photo on Wikipedia / Wikimedia Commons, then 5. a representative photo.
    key = "name:" + k
    hit = cache.get(key)
    if hit:
        if hit.get("file") and Path(hit["file"]).exists():
            return Path(hit["file"]), hit.get("note", "")
        if not hit.get("file") and time.time() - hit.get("checked", 0) < MISS_RETRY_SECONDS:
            return None, ""
    if not web["ok"]:
        return None, ""
    try:
        found, note_prefix = _wikipedia_photo(session, item.name) or _commons_photo(session, item.name), "Photo"
        if not found and (topic := _generic_topic(item)):
            found, note_prefix = _topic_photo(session, topic), "Representative photo"
        if found:
            url, source = found
            path = _download_image(session, url, CACHE_DIR / _safe_filename(item.name))
            note = f"{note_prefix} · {source}"
            cache[key] = {"file": str(path), "note": note, "url": url}
            log(f"Found a web photo for {item.name} ({note}).")
            return path, note
    except (requests.ConnectionError, requests.Timeout) as e:
        # Offline or blocked: skip web lookups for the rest of this run, try again next time.
        web["ok"] = False
        log(f"Web photo lookup unavailable ({e.__class__.__name__}); using photos already saved.")
        return None, ""
    except Exception as e:  # noqa: BLE001 — a bad image file shouldn't stop the PDF
        log(f"Couldn't use the web photo for {item.name}: {e}")
        return None, ""
    cache[key] = {"file": None, "checked": time.time()}
    log(f"No photo found for {item.name}. Add one to Label Maker\\images\\{stem}.jpg")
    return None, ""


# ── PDF ─────────────────────────────────────────────────────────────────────
def _style(name, size, font=None, color=INK, leading=None, **kw):
    return ParagraphStyle(name, fontName=font or FONT, fontSize=size, leading=leading or size * 1.25, textColor=color, **kw)


S_TITLE = _style("title", 22, FONT_BOLD, leading=26)
S_SUB = _style("sub", 9, color=MUTED)
S_HEAD = _style("head", 8, FONT_BOLD, color=colors.white)
S_CELL = _style("cell", 8.5)
S_CELL_B = _style("cellb", 9, FONT_BOLD)
S_CELL_R = _style("cellr", 9, FONT_BOLD, alignment=TA_RIGHT)
S_TINY = _style("tiny", 6.5, color=MUTED)
S_CAT = _style("cat", 6.5, FONT_BOLD, color=ACCENT_DARK)
S_NAME = _style("name", 11.5, FONT_BOLD, leading=13)
S_PRICE = _style("price", 17, FONT_BOLD, leading=19)
S_SPEC = _style("spec", 7.6, leading=9.4)


def _esc(text: str) -> str:
    return str(text).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _money(v) -> str:
    if v is None:
        return "—"
    v = round(float(v), 2)
    s = f"{v:,.2f}" if v % 1 else f"{v:,.0f}"
    # Indian digit grouping: 1,23,456
    whole, _, frac = s.replace(",", "").partition(".")
    if len(whole) > 3:
        head, tail = whole[:-3], whole[-3:]
        head = re.sub(r"(\d)(?=(\d\d)+$)", r"\1,", head)
        whole = f"{head},{tail}"
    return RUPEE + whole + (f".{frac}" if frac else "")


def _photo_flowable(item: Item, w: float, h: float):
    if not item.photo:
        box = Table([[Paragraph("No photo", _style("np", 7, color=MUTED, alignment=1))]], colWidths=[w], rowHeights=[h])
        box.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), PAPER), ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
        return box
    with PILImage.open(item.photo) as im:
        iw, ih = im.size
    scale = min(w / iw, h / ih)
    return Image(str(item.photo), width=iw * scale, height=ih * scale)


def _header_footer(canvas, doc):
    canvas.saveState()
    canvas.setFont(FONT, 7)
    canvas.setFillColor(MUTED)
    canvas.drawString(doc.leftMargin, 10 * mm, f"Zach Store · generated {doc.generated}")
    canvas.drawRightString(A4[0] - doc.rightMargin, 10 * mm, f"Page {doc.page}")
    canvas.setStrokeColor(ACCENT)
    canvas.setLineWidth(2)
    canvas.line(doc.leftMargin, A4[1] - 10 * mm, A4[0] - doc.rightMargin, A4[1] - 10 * mm)
    canvas.restoreState()


def _list_section(items: list[Item], width: float):
    cols = [16 * mm, 22 * mm, 52 * mm, 0, 22 * mm, 13 * mm]
    cols[3] = width - sum(cols)
    head = [Paragraph(t, S_HEAD) for t in ("#", "Photo", "Item", "Specs", "Price", "Stock")]
    rows = [head]
    for i, it in enumerate(items, 1):
        specs = "<br/>".join("· " + _esc(s) for s in it.specs) or "<font color='#605d5d'>Add specs in Item Specs.xlsx</font>"
        stock = "—" if it.stock is None else f"{it.stock:g}"
        stock_style = _style("stk", 9, FONT_BOLD, color=ACCENT_DARK if (it.stock is not None and it.stock <= LOW_STOCK) else INK, alignment=TA_RIGHT)
        rows.append([
            Paragraph(f"{i:02d}", S_CELL_B),
            _photo_flowable(it, 19 * mm, 19 * mm),
            [Paragraph(_esc(it.name), S_CELL_B), Paragraph(_esc(it.type.upper()), S_CAT)],
            Paragraph(specs, S_CELL),
            Paragraph(_money(it.price), S_CELL_R),
            Paragraph(stock, stock_style),
        ])
    t = Table(rows, colWidths=cols, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), INK),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LINEBELOW", (0, 1), (-1, -1), 0.5, RULE),
        ("TOPPADDING", (0, 1), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 1), (-1, -1), 5),
        ("ALIGN", (1, 1), (1, -1), "CENTER"),
    ]))
    return t


def _label(item: Item, w: float, h: float):
    pad = 4 * mm
    photo_w = 30 * mm
    text_w = w - photo_w - pad * 3
    specs = item.specs[:5] or ["Specs not added yet"]
    text = [
        Paragraph(_esc(item.type.upper() or "ITEM"), S_CAT),
        Spacer(1, 1.2 * mm),
        Paragraph(_esc(item.name), S_NAME),
        Spacer(1, 1.6 * mm),
        Paragraph(_money(item.price), S_PRICE),
        Spacer(1, 1.6 * mm),
    ] + [Paragraph("· " + _esc(s), S_SPEC) for s in specs]
    photo = [_photo_flowable(item, photo_w, photo_w)]
    if item.photo_note:
        photo += [Spacer(1, 1 * mm), Paragraph(_esc(item.photo_note), S_TINY)]
    inner = Table([[KeepInFrame(photo_w, h - pad * 2, photo, mode="shrink"),
                    KeepInFrame(text_w, h - pad * 2, text, mode="shrink")]],
                  colWidths=[photo_w + pad, text_w + pad], rowHeights=[h])
    inner.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), pad), ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), pad), ("BOTTOMPADDING", (0, 0), (-1, -1), pad),
    ]))
    return inner


def _labels_section(items: list[Item], width: float):
    per_row, label_h = 2, 64 * mm
    label_w = width / per_row
    cells = [_label(it, label_w, label_h) for it in items]
    while len(cells) % per_row:
        cells.append("")
    grid = [cells[i:i + per_row] for i in range(0, len(cells), per_row)]
    t = Table(grid, colWidths=[label_w] * per_row, rowHeights=[label_h] * len(grid))
    t.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.6, RULE, None, (3, 3)),   # dashed cut lines
        ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    return t


def build_pdf(items: list[Item], out: Path, source_name: str):
    tmp = out.with_name(out.stem + ".tmp.pdf")
    doc = SimpleDocTemplate(str(tmp), pagesize=A4, leftMargin=14 * mm, rightMargin=14 * mm,
                            topMargin=16 * mm, bottomMargin=16 * mm,
                            title="Zach Store — Item Labels", author="Zach Store")
    doc.generated = datetime.now().strftime("%d %b %Y, %H:%M")
    width = A4[0] - doc.leftMargin - doc.rightMargin

    logo = Image(str(LOGO), width=24 * mm, height=24 * mm * 620 / 1030) if LOGO.exists() else Spacer(1, 1)
    in_stock = sum(1 for it in items if (it.stock or 0) > 0)
    title = Table([[logo, [Paragraph("Item list &amp; shelf labels", S_TITLE),
                           Paragraph(f"{len(items)} items · {in_stock} in stock · from {_esc(source_name)} · {doc.generated}", S_SUB)]]],
                  colWidths=[30 * mm, width - 30 * mm])
    title.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0)]))

    story = [title, Spacer(1, 6 * mm), _list_section(items, width), PageBreak(),
             Paragraph("Shelf labels", S_TITLE), Paragraph("Cut along the dashed lines.", S_SUB), Spacer(1, 4 * mm),
             _labels_section(items, width)]
    doc.build(story, onFirstPage=_header_footer, onLaterPages=_header_footer)
    os.replace(tmp, out)   # raises PermissionError if a viewer has the PDF locked


# ── Entry point ─────────────────────────────────────────────────────────────
def generate(inventory: Path, log=print) -> tuple[Path, int]:
    """Build the PDF next to the inventory file. Returns (pdf path, item count)."""
    inventory = Path(inventory)
    items = read_inventory(inventory)
    specs_file = specs_path_for(inventory)
    sync_specs_file(specs_file, items, log)
    specs = read_specs(specs_file)
    for it in items:
        extra_specs, extra_image = specs.get(_key(it.name), ([], ""))
        it.specs = it.specs or extra_specs
        it.image_hint = it.image_hint or extra_image

    IMAGES_DIR.mkdir(parents=True, exist_ok=True)
    cache = _load_cache()
    session, web = _session(), {"ok": True}
    for it in items:
        it.photo, it.photo_note = find_photo(it, session, cache, web, log)
    _save_cache(cache)

    out = inventory.with_name(PDF_NAME)
    build_pdf(items, out, inventory.name)
    return out, len(items)


if __name__ == "__main__":
    import sys
    path, n = generate(Path(sys.argv[1]) if len(sys.argv) > 1 else default_inventory())
    print(f"Wrote {path} ({n} items)")
